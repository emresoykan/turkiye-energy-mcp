"""YTBS İstatistikler sayfası ve günlük Excel raporu ayrıştırıcıları.

Sayfa, 15 dk yük ve kaynak bazlı üretimi gömülü JSON dizileri olarak taşır
(`gunlukYukEgrisiData`, `gunlukUretimEgrisiData`, `santralKaynakListesi`). Geçmiş gün için
PrimeFaces AJAX yanıtı bir JSF partial-response'tur ve CDATA'yı dizi sonunda
`]]><![CDATA[` ile böler; ayrıştırmadan önce bu ayraç kaldırılır.

Günlük Excel raporunda sayfa adları (RaporNNN) anlamsızdır ve başlık satırlarının yeri boş
satırlar nedeniyle değişir; sayfalar ilk satırdaki rapor başlığıyla, tablolar da başlık
hücresinin içeriğiyle bulunur.
"""

from __future__ import annotations

import html
import json
import re
import warnings
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any

import openpyxl

from ..exceptions import EnergyDataError, ErrorCode

SOURCE = "TEİAŞ YTBS"
CDATA_SPLIT = "]]><![CDATA["
_TR_MAP = str.maketrans("çğıöşüÇĞİÖŞÜâÂ", "cgiosuCGIOSUaA")
# Kuruluş × kaynak matrislerindeki ara toplam sütunları; uzun formatta çift sayımı önlemek
# için atlanır.
_SUBTOTALS = {"termik", "termiktoplam", "hidrolik", "hidroliktoplam", "toplam", "toplammw"}


def slug(text: Any) -> str:
    s = str(text).strip().translate(_TR_MAP).lower()
    s = s.replace("(", " ").replace(")", " ").replace("%", " yuzde ")
    return "_".join(s.split())


def source_code(text: Any) -> str:
    """Kaynak adını sayfadaki `santralKaynakListesi` kodlarıyla aynı biçime getirir
    ("Doğal Gaz" -> "dogalgaz", "İthal Kömür" -> "ithalkomur")."""
    return slug(text).replace("_", "")


def _num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).replace(",", "."))


def _tr_date(text: Any) -> date:
    return datetime.strptime(str(text).strip(), "%d.%m.%Y").date()


def _parsing_error(message: str, **details: Any) -> EnergyDataError:
    return EnergyDataError(ErrorCode.PARSING_ERROR, message, source=SOURCE, details=details or None)


# --- İstatistikler sayfası ---


@dataclass(frozen=True, slots=True)
class FormState:
    post_path: str
    view_state: str
    hidden1: str
    year: str
    daily_excel_button: str | None


def _first(pattern: str, text: str, what: str) -> str:
    match = re.search(pattern, text, re.DOTALL)
    if not match:
        raise _parsing_error(f"YTBS sayfasında {what} bulunamadı; sayfa yapısı değişmiş olabilir.")
    return match.group(1)


def parse_form_state(page_html: str) -> FormState:
    buttons = re.findall(r'type="image" src="[^"]*excel\.png[^"]*" name="([^"]+)"', page_html)
    return FormState(
        post_path=html.unescape(_first(r'id="formdash"[^>]*action="([^"]+)"', page_html, "form")),
        view_state=_first(
            r'name="jakarta\.faces\.ViewState"[^>]*value="([^"]+)"', page_html, "ViewState"
        ),
        hidden1=_first(r'name="hidden1" value="([^"]*)"', page_html, "hidden1"),
        year=_first(r'<option value="(\d{4})" selected', page_html, "seçili yıl"),
        daily_excel_button=buttons[0] if buttons else None,
    )


def _js_array(text: str, name: str) -> list[dict[str, Any]] | None:
    match = re.search(rf"var\s+{name}\s*=\s*(\[.*?\]);", text, re.DOTALL)
    if not match:
        return None
    try:
        return json.loads(match.group(1))
    except json.JSONDecodeError as exc:
        raise _parsing_error(f"YTBS {name} JSON olarak okunamadı.") from exc


def _normalize_time(label: str) -> str:
    """ "2026-10-18 24:00" -> "2026-10-19 00:00"; diğer etiketler aynen."""
    day_part, time_part = label.split(" ")
    if time_part == "24:00":
        return (date.fromisoformat(day_part) + timedelta(days=1)).isoformat() + " 00:00"
    return label


def parse_day_series(text: str, day: date) -> list[dict[str, Any]]:
    """Sayfa HTML'i veya AJAX yanıtından 15 dk yük (MW) ve kaynak bazlı üretim (MW) satırları.

    Yük serisi anlık talep değeridir (günlük maksimumu Excel'deki "Maksimum anlık talep" ile
    aynıdır); MWh'e çevrilmemelidir. Henüz yayımlanmamış dilimler 0 gelir, None yapılır.
    """
    text = text.replace(CDATA_SPLIT, "")
    load = _js_array(text, "gunlukYukEgrisiData")
    generation = _js_array(text, "gunlukUretimEgrisiData")
    sources = _js_array(text, "santralKaynakListesi")
    if load is None or generation is None or sources is None:
        raise _parsing_error("YTBS yanıtında 15 dk veri dizileri bulunamadı.")
    if not load or not load[0].get("zaman", "").startswith(day.isoformat()):
        raise EnergyDataError(
            ErrorCode.DATA_NOT_AVAILABLE,
            "YTBS istenen günü döndürmedi.",
            source=SOURCE,
            details={
                "requested": day.isoformat(),
                "returned": load[0].get("zaman") if load else None,
            },
        )

    codes = [s["kod"] for s in sources]
    by_time = {g["zaman"]: g for g in generation}
    rows = []
    for point in load:
        gen = by_time.get(point["zaman"])
        has_gen = bool(gen and gen.get("toplam"))
        row: dict[str, Any] = {
            "tarih": day.isoformat(),
            "zaman": _normalize_time(point["zaman"]),
            "yuk_mw": point["deger"] or None,
        }
        for code in codes:
            row[f"{code}_mw"] = gen.get(code) if has_gen else None
        row["toplam_uretim_mw"] = gen["toplam"] if has_gen else None
        rows.append(row)
    return rows


# --- Günlük Excel raporu ---


def _rows(ws) -> list[list[Any]]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _sheet(wb, title_fragment: str):
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if any(isinstance(v, str) and title_fragment in v for v in first):
            return ws
    return None


def _table(ws, marker: str) -> tuple[list[Any], list[str], list[list[Any]]]:
    rows = _rows(ws)
    key = slug(marker)
    for i, row in enumerate(rows):
        if any(v is not None and slug(v) == key for v in row):
            return row, [slug(h) if h is not None else "" for h in row], rows[i + 1 :]
    raise _parsing_error(f"YTBS raporunda '{marker}' başlığı bulunamadı.", sheet=ws.title)


def _hour_start(day: date, label: str) -> str:
    """Saat-sonu etiketi ("01:00" = 00:00-01:00) için aralık başlangıcı. VARSAYIM: rapor
    saatleri 01:00..24:00 aralığında olduğundan saat-sonu konvansiyonu kabul edilmiştir."""
    hour = int(str(label).split(":")[0])
    start = datetime.combine(day, datetime.min.time()) + timedelta(hours=hour - 1)
    return start.strftime("%Y-%m-%d %H:%M")


def report_date(wb) -> date:
    for row in _rows(wb.worksheets[0])[:3]:
        for value in row:
            if isinstance(value, str):
                try:
                    return _tr_date(value)
                except ValueError:
                    continue
    raise _parsing_error("YTBS raporunun tarihi bulunamadı.")


def _hourly_balance(wb, day: date) -> list[dict[str, Any]]:
    ws = _sheet(wb, "TÜKETİMİNİN SAATLİK GELİŞİMİ")
    if ws is None:
        return []
    _, headers, data = _table(ws, "SAAT")
    out = []
    for r in data:
        if not isinstance(r[0], str) or ":" not in r[0]:
            continue
        rec: dict[str, Any] = {
            "tarih": day.isoformat(),
            "saat": r[0],
            "zaman_baslangic": _hour_start(day, r[0]),
        }
        for h, v in zip(headers[1:], r[1:]):
            if h:
                rec[f"{h}_mwh"] = _num(v)
        out.append(rec)
    return out


def _owner_source_matrix(wb, day: date, title: str, value_name: str) -> list[dict[str, Any]]:
    ws = _sheet(wb, title)
    if ws is None:
        return []
    _, headers, data = _table(ws, "KURULUŞ")
    out = []
    for r in data:
        owner = r[0]
        if not owner or slug(owner).startswith("toplam"):
            continue
        for h, v in zip(headers[1:], r[1:]):
            code = h.replace("_", "")
            if not code or code in _SUBTOTALS:
                continue
            out.append(
                {"tarih": day.isoformat(), "kurulus": owner, "kaynak": code, value_name: _num(v)}
            )
    return out


def _available_capacity(wb, day: date) -> list[dict[str, Any]]:
    ws = _sheet(wb, "EMRE AMADE KAPASİTE")
    if ws is None:
        return []
    raw, _, data = _table(ws, "KAYNAK ADI")
    hours = [(i, h) for i, h in enumerate(raw) if isinstance(h, str) and ":" in h]
    out = []
    for r in data:
        if not r[0]:
            continue
        item = source_code(r[0])
        for i, label in hours:
            out.append(
                {
                    "tarih": day.isoformat(),
                    "saat": label,
                    "zaman_baslangic": _hour_start(day, label),
                    "kalem": item,
                    "deger": _num(r[i]),
                    "birim": r[2],
                }
            )
    return out


def _daily_split(wb, day: date, title: str, key: str) -> list[dict[str, Any]]:
    ws = _sheet(wb, title)
    if ws is None:
        return []
    _, _, data = _table(ws, "AD")
    out = []
    for r in data:
        name = r[1]
        if not name or slug(name) == "toplam":
            continue
        value = source_code(name) if key == "kaynak" else name
        out.append({"tarih": day.isoformat(), key: value, "uretim_mwh": _num(r[2])})
    return out


def _peaks(wb, day: date) -> list[dict[str, Any]]:
    ws = _sheet(wb, "GÜNLÜK TÜRKİYE PUANT RAPORU")
    if ws is None:
        return []
    _, _, data = _table(ws, "ULUSAL PUANT")
    out = []
    for r in data:
        item, period, when, hour, value = (r + [None] * 5)[:5]
        if not item:
            continue
        out.append(
            {
                "tarih": day.isoformat(),
                "kalem": item,
                "donem": period,
                "olay_tarihi": _tr_date(when).isoformat() if when else None,
                "olay_saati": hour,
                "deger": _num(value),
                "birim": "MWh" if "MWh" in item else "MW",
            }
        )
    return out


def _regional_peaks(wb, day: date) -> list[dict[str, Any]]:
    ws = _sheet(wb, "BÖLGESEL PUANT RAPORU")
    if ws is None:
        return []
    _, _, data = _table(ws, "BÖLGE ADI")
    out = []
    for r in data:
        if not r[1]:
            continue
        out.append(
            {
                "tarih": day.isoformat(),
                "bolge": r[1],
                "puant_mw": _num(r[2]),
                "puant_gunu": _tr_date(r[3]).isoformat() if r[3] else None,
                "puant_zamani": r[4],
            }
        )
    return out


def _water(wb, day: date) -> list[dict[str, Any]]:
    """Aylık su durumu (bin m³). Değerlerin dönemi (günlük/kümülatif) doğrulanmadı."""
    ws = _sheet(wb, "SU DURUMU")
    if ws is None:
        return []
    _, headers, data = _table(ws, "GÜN")
    out = []
    for r in data:
        if not isinstance(r[0], (int, float)):
            continue
        rec: dict[str, Any] = {"tarih": day.isoformat(), "ayin_gunu": int(r[0])}
        for h, v in zip(headers[1:], r[1:]):
            if h:
                rec[f"{h}_bin_m3"] = _num(v)
        out.append(rec)
    return out


EXCEL_TABLES = (
    "saatlik_denge",
    "kurulu_guc",
    "santral_sayisi",
    "emre_amade",
    "gunluk_uretim_kaynak",
    "gunluk_uretim_kurulus",
    "puant",
    "bolgesel_puant",
    "su_durumu",
)


def parse_daily_report(content: bytes, expected_day: date | None = None) -> dict[str, list[dict]]:
    """Günlük Excel raporunu tablo adı -> satır listesi sözlüğüne çevirir."""
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        try:
            wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
        except Exception as exc:
            raise _parsing_error("YTBS günlük raporu XLSX olarak açılamadı.") from exc
    day = report_date(wb)
    if expected_day and day != expected_day:
        raise EnergyDataError(
            ErrorCode.DATA_NOT_AVAILABLE,
            "YTBS günlük raporu istenen günü içermiyor.",
            source=SOURCE,
            details={"requested": expected_day.isoformat(), "returned": day.isoformat()},
        )
    return {
        "saatlik_denge": _hourly_balance(wb, day),
        "kurulu_guc": _owner_source_matrix(
            wb, day, "KURULUŞ VE KAYNAKLARA GÖRE KURULU GÜÇ", "kurulu_guc_mw"
        ),
        "santral_sayisi": _owner_source_matrix(
            wb, day, "KURULUŞ VE KAYNAKLARA GÖRE SANTRAL SAYILARI", "santral_sayisi"
        ),
        "emre_amade": _available_capacity(wb, day),
        "gunluk_uretim_kaynak": _daily_split(
            wb, day, "GÜNLÜK ÜRETİMİN KAYNAKLARA GÖRE DAĞILIMI", "kaynak"
        ),
        "gunluk_uretim_kurulus": _daily_split(
            wb, day, "GÜNLÜK ÜRETİMİN KURULUŞLARA GÖRE DAĞILIMI", "kurulus"
        ),
        "puant": _peaks(wb, day),
        "bolgesel_puant": _regional_peaks(wb, day),
        "su_durumu": _water(wb, day),
    }


def balance_gap_mwh(hourly: list[dict[str, Any]]) -> float | None:
    """Saatlik denge: üretim + ithalat - ihracat + deşarj - şarj - tüketim; en büyük |fark|."""
    gaps = []
    for r in hourly:
        try:
            gaps.append(
                abs(
                    r["yenilenemez_uretim_mwh"]
                    + r["yenilenebilir_uretim_mwh"]
                    + r["ithalat_mwh"]
                    - r["ihracat_mwh"]
                    + r["depolama_desarj_mwh"]
                    - r["depolama_sarj_mwh"]
                    - r["tuketim_mwh"]
                )
            )
        except (KeyError, TypeError):
            return None
    return max(gaps) if gaps else None
