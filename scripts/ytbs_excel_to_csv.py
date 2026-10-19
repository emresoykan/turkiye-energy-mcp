"""TEİAŞ YTBS günlük Excel raporlarını (ytbs_gunluk_YYYY-MM-DD.xlsx) düzenli CSV'lere çevirir.

Excel'ler `scripts/ytbs_fetch.ps1 -Excel` ile indirilir. Klasördeki tüm günlük raporlar
okunur ve her tablo için tek bir "uzun" (tidy) CSV yeniden yazılır; aynı gün birden fazla
dosyada varsa en son değiştirilen dosya kullanılır (YTBS verileri geriye dönük revize olur).

Kullanım:
    python ytbs_excel_to_csv.py                  # varsayılan: ~/Desktop/ytbs
    python ytbs_excel_to_csv.py C:\\yol\\klasor
    python ytbs_excel_to_csv.py dosya1.xlsx dosya2.xlsx --cikti C:\\yol\\csv

Notlar:
- Veriler TEİAŞ'a göre geçici işletme verisidir; geriye dönük düzeltmeye açıktır ve lisanssız
  üretim ay sonuna kadar tahminidir.
- Saatlik tablolarda "01:00" ... "24:00" saat-sonu (hour-ending) etiketi olarak yorumlanır;
  `zaman_baslangic` bu VARSAYIMLA türetilir (ör. "01:00" -> 00:00-01:00 aralığı).
- Aylık kümülatif sayfalar (ay içi günlük gelişim, aylık dağılımlar) ve su durumu raporu
  (dönemi doğrulanmadı) çıktıya alınmaz.

Gereksinim: pip install openpyxl
"""

from __future__ import annotations

import argparse
import csv
import sys
import warnings
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl gerekli: pip install openpyxl")

warnings.filterwarnings("ignore", message="Workbook contains no default style")

TR_MAP = str.maketrans("çğıöşüÇĞİÖŞÜâÂ", "cgiosuCGIOSUaA")
# Ara toplam sütunları; uzun formatta tekrar sayımı önlemek için atlanır.
SUBTOTAL_COLUMNS = {"termik", "termik_toplam", "hidrolik", "hidrolik_toplam", "toplam", "toplam_mw"}
BALANCE_TOLERANCE_MWH = 1.0


def slug(text) -> str:
    s = str(text).strip().translate(TR_MAP).lower()
    s = s.replace("(", " ").replace(")", " ").replace("%", " yuzde ")
    return "_".join(s.split())


def num(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).replace(",", "."))


def parse_tr_date(text) -> date:
    return datetime.strptime(str(text).strip(), "%d.%m.%Y").date()


def hour_start(day: date, label: str) -> str:
    hour = int(str(label).split(":")[0])
    return (datetime.combine(day, datetime.min.time()) + timedelta(hours=hour - 1)).strftime("%Y-%m-%d %H:%M")


def rows_of(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def find_sheet(wb, title_fragment: str):
    """Sayfayı adıyla değil, ilk satırdaki rapor başlığıyla bulur (sayfa adları RaporNNN)."""
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if any(isinstance(v, str) and title_fragment in v for v in first):
            return ws
    return None


def table(ws, marker: str):
    """Başlık satırını (içinde `marker` hücresi olan ilk satır) bulur; araya boş satır
    girebildiği için satır numarası sabit değildir. (ham başlıklar, slug başlıklar, veri) döndürür."""
    rows = rows_of(ws)
    key = slug(marker)
    for i, row in enumerate(rows):
        if any(v is not None and slug(v) == key for v in row):
            return row, [slug(h) if h is not None else "" for h in row], rows[i + 1:]
    raise ValueError(f"'{marker}' başlık satırı bulunamadı: {ws.title}")


def report_date(ws) -> date:
    for row in rows_of(ws)[:3]:
        for v in row:
            if isinstance(v, str):
                try:
                    return parse_tr_date(v)
                except ValueError:
                    continue
    raise ValueError(f"Rapor tarihi bulunamadı: {ws.title}")


# --- Tablo ayrıştırıcıları: her biri sözlük listesi döndürür ---

def parse_hourly_balance(wb, day, src):
    ws = find_sheet(wb, "TÜKETİMİNİN SAATLİK GELİŞİMİ")
    if ws is None:
        return []
    _, headers, data = table(ws, "SAAT")
    out = []
    for r in data:
        label = r[0]
        if not isinstance(label, str) or ":" not in label:
            continue
        rec = {"tarih": day.isoformat(), "saat": label, "zaman_baslangic": hour_start(day, label)}
        for h, v in zip(headers[1:], r[1:]):
            if h:
                rec[f"{h}_mwh"] = num(v)
        rec["kaynak_dosya"] = src
        out.append(rec)
    return out


def parse_matrix(wb, day, src, title, row_key, value_name):
    """Kuruluş × kaynak matrislerini (kurulu güç, santral sayısı) uzun formata çevirir."""
    ws = find_sheet(wb, title)
    if ws is None:
        return []
    _, headers, data = table(ws, "KURULUŞ")
    out = []
    for r in data:
        name = r[0]
        if not name or slug(name).startswith("toplam"):
            continue
        for h, v in zip(headers[1:], r[1:]):
            if not h or h in SUBTOTAL_COLUMNS:
                continue
            out.append({"tarih": day.isoformat(), row_key: name, "kaynak": h,
                        value_name: num(v), "kaynak_dosya": src})
    return out


def parse_available_capacity(wb, day, src):
    ws = find_sheet(wb, "EMRE AMADE KAPASİTE")
    if ws is None:
        return []
    raw, _, data = table(ws, "KAYNAK ADI")
    hour_cols = [(i, h) for i, h in enumerate(raw) if isinstance(h, str) and ":" in h]
    out = []
    for r in data:
        item = r[0]
        if not item:
            continue
        unit = r[2]
        for i, label in hour_cols:
            out.append({"tarih": day.isoformat(), "saat": label, "zaman_baslangic": hour_start(day, label),
                        "kalem": item, "deger": num(r[i]), "birim": unit, "kaynak_dosya": src})
    return out


def parse_daily_split(wb, day, src, title, key):
    ws = find_sheet(wb, title)
    if ws is None:
        return []
    _, _, data = table(ws, "AD")
    out = []
    for r in data:
        name, value = r[1], r[2]
        if not name or slug(name) == "toplam":
            continue
        out.append({"tarih": day.isoformat(), key: name, "uretim_mwh": num(value), "kaynak_dosya": src})
    return out


def parse_peak(wb, day, src):
    ws = find_sheet(wb, "GÜNLÜK TÜRKİYE PUANT RAPORU")
    if ws is None:
        return []
    _, _, data = table(ws, "ULUSAL PUANT")
    out = []
    for r in data:
        item, period, when, hour, value = (r + [None] * 5)[:5]
        if not item:
            continue
        unit = "MWh" if "MWh" in item else "MW"
        out.append({"rapor_tarihi": day.isoformat(), "kalem": item, "donem": period,
                    "tarih": parse_tr_date(when).isoformat() if when else None,
                    "saat": hour, "deger": num(value), "birim": unit, "kaynak_dosya": src})
    return out


def parse_regional_peak(wb, day, src):
    ws = find_sheet(wb, "BÖLGESEL PUANT RAPORU")
    if ws is None:
        return []
    _, _, data = table(ws, "BÖLGE ADI")
    out = []
    for r in data:
        if not r[1]:
            continue
        out.append({"tarih": day.isoformat(), "bolge": r[1], "puant_mw": num(r[2]),
                    "puant_gunu": parse_tr_date(r[3]).isoformat() if r[3] else None,
                    "puant_zamani": r[4], "kaynak_dosya": src})
    return out


OUTPUTS = {
    "ytbs_saatlik_denge.csv": parse_hourly_balance,
    "ytbs_kurulu_guc.csv": lambda wb, d, s: parse_matrix(
        wb, d, s, "KURULUŞ VE KAYNAKLARA GÖRE KURULU GÜÇ", "kurulus", "kurulu_guc_mw"),
    "ytbs_santral_sayisi.csv": lambda wb, d, s: parse_matrix(
        wb, d, s, "KURULUŞ VE KAYNAKLARA GÖRE SANTRAL SAYILARI", "kurulus", "santral_sayisi"),
    "ytbs_emre_amade.csv": parse_available_capacity,
    "ytbs_gunluk_uretim_kaynak.csv": lambda wb, d, s: parse_daily_split(
        wb, d, s, "GÜNLÜK ÜRETİMİN KAYNAKLARA GÖRE DAĞILIMI", "kaynak"),
    "ytbs_gunluk_uretim_kurulus.csv": lambda wb, d, s: parse_daily_split(
        wb, d, s, "GÜNLÜK ÜRETİMİN KURULUŞLARA GÖRE DAĞILIMI", "kurulus"),
    "ytbs_puant.csv": parse_peak,
    "ytbs_bolgesel_puant.csv": parse_regional_peak,
}


def check_balance(rows, day) -> list[str]:
    """Üretim + ithalat - ihracat + deşarj - şarj = tüketim kontrolü (saatlik)."""
    issues = []
    worst = 0.0
    for r in rows:
        try:
            lhs = (r["yenilenemez_uretim_mwh"] + r["yenilenebilir_uretim_mwh"] + r["ithalat_mwh"]
                   - r["ihracat_mwh"] + r["depolama_desarj_mwh"] - r["depolama_sarj_mwh"])
            worst = max(worst, abs(lhs - r["tuketim_mwh"]))
        except (KeyError, TypeError):
            issues.append(f"{day}: saatlik denge sütunları beklenen adlarda değil, kontrol atlandı")
            return issues
    if worst > BALANCE_TOLERANCE_MWH:
        issues.append(f"{day}: saatlik enerji dengesi farkı {worst:.2f} MWh (> {BALANCE_TOLERANCE_MWH})")
    return issues


def write_csv(path: Path, rows: list[dict]):
    if not rows:
        return
    fields = list(dict.fromkeys(k for r in rows for k in r))
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in fields})


def collect_files(inputs: list[str]) -> list[Path]:
    paths = []
    for item in inputs:
        p = Path(item).expanduser()
        if p.is_dir():
            paths.extend(sorted(p.glob("ytbs_gunluk_*.xlsx")))
        elif p.suffix.lower() == ".xlsx":
            paths.append(p)
    return paths


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("girdi", nargs="*", default=[str(Path.home() / "Desktop" / "ytbs")],
                    help="Klasör(ler) veya .xlsx dosya(lar)")
    ap.add_argument("--cikti", help="CSV klasörü (varsayılan: ilk girdinin klasörü/csv)")
    args = ap.parse_args(argv)

    files = collect_files(args.girdi)
    if not files:
        print("Hiç ytbs_gunluk_*.xlsx bulunamadı.", file=sys.stderr)
        return 1
    first = Path(args.girdi[0]).expanduser()
    out_dir = Path(args.cikti) if args.cikti else (first if first.is_dir() else first.parent) / "csv"
    out_dir.mkdir(parents=True, exist_ok=True)

    # Aynı rapor tarihi için en son değiştirilen dosyayı seç.
    by_day: dict[date, tuple[float, Path, object]] = {}
    issues: list[str] = []
    for f in files:
        try:
            wb = openpyxl.load_workbook(f, read_only=False, data_only=True)
            day = report_date(wb.worksheets[0])
        except Exception as exc:  # bozuk/HTML dosyası vb.
            issues.append(f"{f.name}: okunamadı ({exc})")
            continue
        mtime = f.stat().st_mtime
        if day not in by_day or mtime > by_day[day][0]:
            by_day[day] = (mtime, f, wb)

    results = {name: [] for name in OUTPUTS}
    for day in sorted(by_day):
        _, f, wb = by_day[day]
        for name, parser in OUTPUTS.items():
            try:
                results[name].extend(parser(wb, day, f.name))
            except Exception as exc:
                issues.append(f"{f.name} / {name}: {exc}")
        issues.extend(check_balance([r for r in results["ytbs_saatlik_denge.csv"] if r["tarih"] == day.isoformat()], day))
        hourly = [r for r in results["ytbs_saatlik_denge.csv"] if r["tarih"] == day.isoformat()]
        if hourly:
            uretim = sum((r.get("yenilenemez_uretim_mwh") or 0) + (r.get("yenilenebilir_uretim_mwh") or 0) for r in hourly)
            tuketim = sum(r.get("tuketim_mwh") or 0 for r in hourly)
            print(f"{day}: {len(hourly)} saat | üretim {uretim:,.0f} MWh | tüketim {tuketim:,.0f} MWh | {f.name}")

    for name, rows in results.items():
        write_csv(out_dir / name, rows)
        print(f"  {name}: {len(rows)} satır")
    print(f"CSV klasörü: {out_dir}")
    for msg in issues:
        print(f"UYARI: {msg}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
